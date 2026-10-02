#!/usr/bin/env python3
"""Round 1 recovery closeout, step 2 - READ-ONLY classification of the 5 returned workbooks.

Classifies every response row in reports/partner_data_recovery/inputs/Round 1_run/ against the CURRENT
tracker and Round 1 data, and proposes an action per row. Writes nothing to the tracker - the apply step is
separate and needs Jack's OK.

Why not the existing verify_data_recovery_response.py / review_recovery_response.py:
  - verify's EXPECTED_HEADERS for "Non-IDP Duplicates" predates the 2026-09-21 claim/dropdown redesign
    (expects "CONFIRMED Genuine Interview ID"; the 22 Sep workbooks carry "CONFIRMED Point ID"), so it would
    reject that whole sheet on the header check, and review_ has no Non-IDP Duplicates step at all.
  - review_idp_listing_duplicates() resolves an "idp_listing_duplicate::<uuid>" tracker row, but since the
    redesign most IDP Listing Duplicates rows are confirmed_deletion/duplicate_point rows - approving them
    that way would leave the real duplicate_point row pending.
  - verify's IDP listing-number check reads idp_real_listing_pools.csv, last exported 11 Sep; the HH
    listing export is from 29 Sep. Pools are rebuilt here from the export with real_hh_listing.R's own
    rule (latest usable submission per cluster, primary_list + reserve_list).
The validation RULES themselves are carried over from verify_ (format, same-cluster, in-pool, not reused).

Only confirmed_deletion rows can ever exclude an interview from Achieved (build_confirmed_deletions_overlay.R
filters issue_type == "confirmed_deletion"; settled = terminal status AND recovery_type NA). So a recovered
duplicate is recorded as status=confirmed + recovery_type=false_positive and stays in Achieved.

    python -B _working_files/scripts/round1_ingest_classify.py
"""
import csv
import os
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[2]
RET_DIR = REPO / "reports/partner_data_recovery/inputs/Round 1_run"
ORIG_DIR = REPO / "reports/partner_data_recovery/outputs"
OUT_DIR = REPO / "reports/partner_data_recovery/outputs/_round1_closeout"
TRACKER = REPO / "reports/partner_data_recovery/scripts/recovery_issue_tracker.csv"
SUBS = REPO / "data/real_submissions.csv"
MEMBERS = REPO / "data/ROUND1_MEMBERSHIP.csv"
FRAME_DIR = REPO / "input_data/sampling_frame"
HH_LISTING = REPO / "cleaning/MSNA_Data_Cleaning/Kobo Downloads/hh_listing_tool/hh_listing.xlsx"

TERMINAL = {"confirmed", "contested"}
NO_APPEAL_NOTE = "No action needed -- confirmed per validated assessment methodology, not open to contest."
CONTESTED_CLOSED_NOTE = "No action needed -- your team has already contested this interview, and the MSNA team has it on record."
POINT_RE = re.compile(r"^non_idp_(NG\d{6})_([0-9A-Za-z]+(?:_supp\d+)?)_(HH|R)\d+$")

# returned file -> (org_id, original outgoing workbook path)
WORKBOOKS = {
    "COOPI_data_recovery_workbook_2026-09-22.xlsx": ("coopi", "COOPI/COOPI_data_recovery_workbook_2026-09-22.xlsx"),
    "IMC_data_recovery_workbook_2026-09-22_JD.xlsx": ("imc", "IMC/IMC_data_recovery_workbook_2026-09-22.xlsx"),
    "MALTESER_data_recovery_workbook_ Updated_Sept_2026.xlsx": ("malteser", "MALTESER/MALTESER_data_recovery_workbook_2026-09-22.xlsx"),
    "NRC_data_recovery_workbook_2026-09-28(1).xlsx": ("nrc", "NRC/NRC_data_recovery_workbook_2026-09-22.xlsx"),
    "PLAN_data_recovery_workbook_2026-09-22.xlsx": ("plan", "PLAN/PLAN_data_recovery_workbook_2026-09-22.xlsx"),
}
SHEETS = ["Confirmed Deletions", "Non-IDP Duplicates", "IDP Listing Duplicates", "Missing HH Listings", "Other Issues"]


def s(v):
    return "" if v is None else str(v).strip()


def read_sheet(path, sheet):
    """(header, rows-as-dicts) - header = first row within the top 6 holding 'Interview ID'/'Cluster/Site ID'."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        return None, []
    ws = wb[sheet]
    ws.reset_dimensions()  # some workbooks carry a stale A1:A1 dimension tag
    rows = list(ws.iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows[:6]) if r and ("Interview ID" in r or "Cluster/Site ID" in r)), None)
    if hi is None:
        return None, []
    header = [s(h) for h in rows[hi]]
    out = []
    for r in rows[hi + 1:]:
        r = list(r or []) + [None] * (len(header) - len(r or []))
        if not any(v not in (None, "") for v in r):
            continue
        out.append({h: r[i] for i, h in enumerate(header) if h})
    return header, out


def latest_frame(kind):
    pat = re.compile(rf"^NGA_MSNA_2026_stage2_sampling_frame_v(\d+)_{kind}\.csv$")
    hits = [(int(m.group(1)), p) for p in FRAME_DIR.iterdir() if (m := pat.match(p.name))]
    return max(hits)[1]


def build_listing_pools():
    """real_hh_listing.R's compute_real_avail_pools(): latest submission per cluster with a non-blank
    primary_list; pool = sorted unique ints of primary_list + reserve_list."""
    wb = openpyxl.load_workbook(HH_LISTING, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    ws.reset_dimensions()
    it = ws.iter_rows(values_only=True)
    hdr = [s(h) for h in next(it)]
    ix = {h: i for i, h in enumerate(hdr)}
    best = {}
    for r in it:
        r = list(r) + [None] * (len(hdr) - len(r))
        cl, pl = s(r[ix["cluster_select"]]), s(r[ix["primary_list"]])
        if not cl or not pl:
            continue
        t = s(r[ix["_submission_time"]])
        if cl not in best or t > best[cl][0]:
            best[cl] = (t, pl, s(r[ix["reserve_list"]]))
    pools = {}
    for cl, (_, pl, rl) in best.items():
        nums = set()
        for tok in (pl + " " + rl).split():
            try:
                nums.add(int(float(tok)))
            except ValueError:
                pass
        pools[cl] = nums
    return pools


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    members = {r["submission_uuid"] for r in csv.DictReader(open(MEMBERS, encoding="utf-8-sig"))}
    tracker = list(csv.DictReader(open(TRACKER, encoding="utf-8-sig")))
    by_id = {t["issue_id"]: t for t in tracker}
    subs = {r["submission_uuid"]: r for r in csv.DictReader(open(SUBS, encoding="utf-8-sig"))}

    # non-IDP point claims, Round 1, for "already claimed by someone else" (IDP listing-number claims are
    # built further down as post_claims, with every proposed correction applied)
    point_claims = defaultdict(set)
    for u, r in subs.items():
        if s(r["non_idp_point_id"]) not in ("", "NA"):
            point_claims[s(r["non_idp_point_id"])].add(u)

    frame = {}
    with open(latest_frame("FULL"), encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            if r["pop_type"] == "non_idp":
                frame[r["survey_id"]] = r["cluster_id"]
    pools = build_listing_pools()
    print(f"loaded: {len(members)} Round 1 uuids | {len(tracker)} tracker rows | {len(frame)} non-IDP frame points | "
          f"{len(pools)} clusters with a real HH listing")

    # IDP Listing Duplicates lists BOTH sides of a shared listing number (the flagged later claimant AND its
    # unflagged first claimant), and partners corrected both - so a correction has to be judged against the
    # cluster's claims AFTER every proposed correction in that cluster is applied, not against today's claims
    # one row at a time (A #5->#7 frees #5 for B). Pre-pass: collect every valid-looking proposed IDP number.
    proposed_idp = {}  # uuid -> new number
    for fname in WORKBOOKS:
        _, rows = read_sheet(RET_DIR / fname, "IDP Listing Duplicates")
        for r in rows:
            u, raw = s(r.get("Interview ID")), r.get("CONFIRMED Listing Number")
            if u in subs and raw not in (None, ""):
                try:
                    proposed_idp[u] = int(float(raw))
                except (TypeError, ValueError):
                    pass
    post_claims = defaultdict(set)  # (cluster, number) -> uuids, with every proposal applied
    for u, r in subs.items():
        clu = s(r["matched_cluster_id"])
        if clu in ("", "NA") or r["pop_type"] != "idp":
            continue
        if u in proposed_idp:
            post_claims[(clu, proposed_idp[u])].add(u)
        elif s(r["idp_hh_number_from_listing"]) not in ("", "NA"):
            try:
                post_claims[(clu, int(float(r["idp_hh_number_from_listing"])))].add(u)
            except ValueError:
                pass

    out_rows = []
    used_points = {}  # Non-IDP proposals across ALL workbooks, to catch the same point used twice

    def add(org, fname, sheet, rownum, key, issue_id, partner_answer, cls, action, detail):
        t = by_id.get(issue_id, {})
        out_rows.append({
            "org_id": org, "file": fname, "sheet": sheet, "row": rownum, "key": key, "issue_id": issue_id,
            "tracker_status": t.get("status", "NOT_IN_TRACKER"), "deletion_reason": t.get("deletion_reason", ""),
            "partner_answer": partner_answer, "classification": cls, "proposed_action": action, "detail": detail,
        })

    for fname, (org, orig_rel) in WORKBOOKS.items():
        ret_path, orig_path = RET_DIR / fname, ORIG_DIR / orig_rel
        for sheet in SHEETS:
            hdr, rows = read_sheet(ret_path, sheet)
            if hdr is None:
                continue
            ohdr, orows = read_sheet(orig_path, sheet) if orig_path.exists() else (None, [])
            idcol = "Interview ID" if "Interview ID" in hdr else "Cluster/Site ID"
            if ohdr is not None and [s(r.get(idcol)) for r in rows] != [s(r.get(idcol)) for r in orows]:
                add(org, fname, sheet, None, "", "", "", "alignment_error", "none",
                    "row set/order differs from the 22 Sep outgoing workbook - review before trusting positions")

            for i, r in enumerate(rows, start=2):
                if sheet == "Missing HH Listings":
                    cl = s(r.get("Cluster/Site ID"))
                    iid = f"missing_hh_listing::{cl}"
                    ans = s(r.get("Household Listing Now Submitted? (Yes/No)"))
                    t = by_id.get(iid)
                    n_lm = sum(1 for x in tracker if x["deletion_reason"] == "listing_missing" and x["cluster_id"] == cl
                               and x["status"] not in TERMINAL)
                    if t is None and n_lm:
                        add(org, fname, sheet, i, cl, "", ans, "defer_q4", "q4_household_number_check",
                            f"{n_lm} open interview-level listing_missing row(s); cluster {'HAS' if cl in pools else 'has NO'} real listing now")
                    elif t is None:
                        add(org, fname, sheet, i, cl, iid, ans, "not_in_tracker", "none", "")
                    elif t["status"] in TERMINAL:
                        add(org, fname, sheet, i, cl, iid, ans, "already_settled", "none", "")
                    elif cl in pools:
                        add(org, fname, sheet, i, cl, iid, ans, "apply", "close_listing_received",
                            f"cluster now has a real listing in the HH listing export ({len(pools[cl])} drawn numbers)")
                    elif ans.lower() == "yes":
                        add(org, fname, sheet, i, cl, iid, ans, "unverifiable", "fallback",
                            "partner says submitted, but the cluster is not in the 29 Sep HH listing export")
                    else:
                        add(org, fname, sheet, i, cl, iid, ans, "no_response", "fallback", "")
                    continue

                uuid = s(r.get("Interview ID"))
                if uuid and uuid not in members:
                    add(org, fname, sheet, i, uuid, "", "", "not_round1", "none", "Interview ID not in ROUND1_MEMBERSHIP")
                    continue
                cd_id = f"confirmed_deletion::{uuid}"

                if sheet == "Confirmed Deletions":
                    ans = s(r.get("Contest This? (Yes/No)"))
                    expl = s(r.get("If Yes, Explain"))
                    t = by_id.get(cd_id)
                    if t is None:
                        add(org, fname, sheet, i, uuid, cd_id, ans, "not_in_tracker", "none", "")
                    elif t["status"] in TERMINAL:
                        note = "Q7: contest ignored, under-20 deletion stands" if ans.lower().startswith("yes") and t["deletion_reason"] == "duration_under_20" else ""
                        add(org, fname, sheet, i, uuid, cd_id, ans[:40], "already_settled", "none", note)
                    elif ans in (NO_APPEAL_NOTE, CONTESTED_CLOSED_NOTE):
                        add(org, fname, sheet, i, uuid, cd_id, "(pre-filled note)", "anomaly", "none",
                            "no-appeal/closed note on a row the tracker still has open")
                    elif ans.lower().startswith("yes"):
                        add(org, fname, sheet, i, uuid, cd_id, ans, "judgment", "contest_review",
                            f"{t['deletion_reason']}: {expl[:200]}")
                    elif ans.lower() in ("no", ""):
                        add(org, fname, sheet, i, uuid, cd_id, ans or "(blank)", "apply" if ans else "no_response",
                            "confirm_deletion_partner" if ans else "fallback", t["deletion_reason"])
                    else:
                        add(org, fname, sheet, i, uuid, cd_id, ans, "judgment", "unclear_answer", t["deletion_reason"])
                    continue

                if sheet == "Non-IDP Duplicates":
                    ans = s(r.get("CONFIRMED Point ID"))
                    claimed = s(r.get("Point ID Claimed"))
                    cluster = s(r.get("Cluster ID"))
                    t = by_id.get(cd_id)
                    if t is None:
                        add(org, fname, sheet, i, uuid, cd_id, ans, "not_in_tracker", "none", f"claimed={claimed}")
                        continue
                    if t["status"] in TERMINAL:
                        add(org, fname, sheet, i, uuid, cd_id, ans, "already_settled", "none", f"claimed={claimed}")
                        continue
                    if not ans:
                        add(org, fname, sheet, i, uuid, cd_id, "", "no_response", "fallback", f"claimed={claimed}")
                        continue
                    problems = []
                    if not POINT_RE.match(ans):
                        problems.append("not a recognisable non-IDP point id")
                    elif ans not in frame:
                        problems.append("point id not in the sampling frame")
                    elif cluster and frame[ans] != cluster:
                        problems.append(f"point is in cluster {frame[ans]}, interview is in {cluster}")
                    if ans == claimed:
                        problems.append("same as the disputed claimed point (partner insists on the original)")
                    others = point_claims.get(ans, set()) - {uuid}
                    if others:
                        problems.append(f"already claimed by {len(others)} other Round 1 interview(s)")
                    if ans in used_points:
                        problems.append(f"same point proposed for another flagged interview ({used_points[ans]})")
                    used_points.setdefault(ans, uuid)
                    if problems:
                        add(org, fname, sheet, i, uuid, cd_id, ans, "judgment", "point_review", f"claimed={claimed}; " + "; ".join(problems))
                    else:
                        add(org, fname, sheet, i, uuid, cd_id, ans, "apply", "recover_corrected_point", f"{claimed} -> {ans}")
                    continue

                if sheet == "IDP Listing Duplicates":
                    raw = r.get("CONFIRMED Listing Number")
                    recorded = s(r.get("Listing Number Recorded")).split(".")[0]
                    cluster = s(r.get("Cluster/Site ID"))
                    iid = cd_id if cd_id in by_id else f"idp_listing_duplicate::{uuid}"
                    t = by_id.get(iid)
                    flagged = t is not None
                    if flagged and t["status"] in TERMINAL:
                        add(org, fname, sheet, i, uuid, iid, s(raw), "already_settled", "none", f"recorded={recorded}")
                        continue
                    if raw in (None, ""):
                        if flagged:
                            add(org, fname, sheet, i, uuid, iid, "", "no_response", "fallback", f"recorded={recorded}")
                        else:
                            add(org, fname, sheet, i, uuid, "", "", "no_change", "none", "unflagged first claimant, no answer")
                        continue
                    try:
                        n = int(float(raw))
                    except (TypeError, ValueError):
                        add(org, fname, sheet, i, uuid, iid if flagged else "", s(raw), "judgment", "listing_review",
                            f"not a whole number; recorded={recorded}; {'flagged' if flagged else 'unflagged first claimant'}")
                        continue
                    problems = []
                    pool = pools.get(cluster)
                    if pool is None:
                        problems.append("cluster has no real HH listing to check the number against")
                    elif n not in pool:
                        problems.append(f"number not in the cluster's drawn listing pool ({len(pool)} numbers)")
                    others = post_claims.get((cluster, n), set()) - {uuid}
                    if others:
                        problems.append(f"still shared with {len(others)} other interview(s) after every proposed correction")
                    who = "flagged" if flagged else "unflagged first claimant"
                    if problems:
                        add(org, fname, sheet, i, uuid, iid if flagged else "", s(n), "judgment", "listing_review",
                            f"{who}; recorded={recorded}; " + "; ".join(problems))
                    elif not flagged:
                        if s(n) == recorded:
                            add(org, fname, sheet, i, uuid, "", s(n), "no_change", "none", "unflagged first claimant, confirms recorded number")
                        else:
                            add(org, fname, sheet, i, uuid, "", s(n), "apply", "correct_listing_log_only",
                                f"unflagged first claimant {recorded} -> {n} (no tracker issue; cleaning-log correction only)")
                    else:
                        action = "recover_keep_listing" if s(n) == recorded else "recover_corrected_listing"
                        add(org, fname, sheet, i, uuid, iid, s(n), "apply", action,
                            f"{recorded} -> {n}" if action == "recover_corrected_listing" else f"keeps {n}; sibling moved off it")
                    continue

                if sheet == "Other Issues":
                    t = by_id.get(cd_id)
                    fields = {k: s(v) for k, v in r.items() if k not in ("Interview ID", "Enumerator ID", "State", "LGA", "Ward", "Date of Submission")}
                    if t is None:
                        add(org, fname, sheet, i, uuid, cd_id, "", "not_in_tracker", "none", "")
                    elif t["status"] in TERMINAL:
                        add(org, fname, sheet, i, uuid, cd_id, "", "already_settled", "none", "")
                    else:
                        add(org, fname, sheet, i, uuid, cd_id, "", "judgment", "other_issue_review",
                            "; ".join(f"{k}={v}" for k, v in fields.items() if v)[:300])

    out_csv = OUT_DIR / "round1_returned_workbooks_classified.csv"
    with open(out_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(out_rows[0].keys()))
        w.writeheader()
        w.writerows(out_rows)

    print(f"\nwrote {out_csv} ({len(out_rows)} rows)\n")
    print("By partner / sheet / classification / proposed action:")
    c = Counter((r["org_id"], r["sheet"], r["classification"], r["proposed_action"]) for r in out_rows)
    for k, v in sorted(c.items()):
        print(f"  {k[0]:9s} {k[1]:24s} {k[2]:16s} {k[3]:28s} {v}")
    print("\nTotals by classification:", dict(Counter(r["classification"] for r in out_rows)))
    print("Judgment items by type:", dict(Counter(r["proposed_action"] for r in out_rows if r["classification"] == "judgment")))


if __name__ == "__main__":
    main()
