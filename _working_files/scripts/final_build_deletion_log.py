#!/usr/bin/env python3
"""FINAL deletion & correction log (Jack, 16 Oct 2026, decision #66): ONE cumulative log over the WHOLE final dataset,
in the data officer's master-cleaning-log format, superseding the Round 1 log of 2 Oct
(outputs/_round1_closeout/MSNA_N-WEC_2026_Round1_deletion_log_2026-10-02.*). Built from round1_build_deletion_log.py,
which stays as the record of what was sent for Round 1 - same 29 columns, change_type/status vocabularies,
check_binding convention and row wording, so a Round 1 row reads the same in both logs. Differences:
  scope           every submission in data/real_submissions.csv (the final refresh), not only the 28,046 of Round 1
  deletion_round  remove_survey rows: "Round 1" = already removed in the Round 1 log sent to the data officer;
                  "Round 2" = removed since - new interviews, and Round 1 interviews removed after the Round 1
                  submission (e.g. the 27 short interviews found on 4 Oct, whose audit files arrived late)
  interview_round every row: "Round 1" = among the 28,046 Round 1 submissions (data/ROUND1_MEMBERSHIP.csv), else "Round 2"
  identity        Round 1 interviews keep the identity columns of the Round 1 log verbatim (the data officer's dataset
                  before the Round 1 corrections; frame values as of 2 Oct); later interviews take them from the final
                  data (no corrections apply) and the current FULL frame, falling back to the archived v14 frame for any
                  cluster no longer in it
  exceptions      written to <stem>_EXCEPTIONS.csv: (1) Round 1 removals since reversed; (2) Achieved interviews with
                  no audit-trail duration (#68 - Jack: wait for the full audit, no placeholder) - except an incomplete
                  audit trail whose minimum (final_audit_consistency.py's lower-bounds file) already reaches 19.95
                  minutes: Jack, 16 Oct, it passes, and gets a duration_audit_incomplete kept row stating the minimum;
                  (3) tracker rows still open; (4) duration removals the audit trail does not support. (2)-(4) stop the
                  build unless --preview.

Every row is justified (Jack, 2 Oct: "make sure ... we're putting reasons/justifications next to each one so it can
also be traced and defended"): issue = what the check found, Data_Feedback = the evidence, FO_Comments = the decision
and why (the master log's own "Survey removed. ..." / "Changed from X to Y. ..." / "No action. ..." style),
response_file = where the decision came from; then deletion_round, interview_round, decision_basis, decided_by,
tracker_issue_id, review_reminder and already_in_do_master_log.

Rows:
  remove_survey   one per settled deletion (the CONFIRMED overlay), labelled by its REAL basis - the tracker's
                  deletion_reason is not always it: one confirmed_deletion row per uuid means a later automatic
                  duration/no_consent confirmation can land on whatever reason was registered first. Rule order is the
                  data officer's own priority: no_consent, then the duration rule (any interview under 19.95 minutes by
                  its audit trail is a duration removal, whatever else it was flagged for), then the rest
  change_response one per (interview, variable) from outputs/_round1_closeout/round1_final_corrections.csv - the SAME
                  file the dashboard's prep overlay reads (data/ROUND1_CORRECTIONS.csv) - for Round 1 interviews that
                  are still kept; no later corrections exist (the final round changed no identifiers)
  no_action       one per flag resolved by keeping the interview with no data change, plus one per interview in a
                  cluster that never had a household listing (IOM DTM site as reference)
Completeness is asserted: every tracker row with a uuid in the final data is referenced by at least one log row, or
its interview is removed.

    python -B _working_files/scripts/final_build_deletion_log.py [--preview] [--out <dir>]
"""
import csv
import datetime
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

REPO = Path(__file__).resolve().parents[2]
CL = REPO / "reports/partner_data_recovery/outputs/_round1_closeout"
OUT_DIR = REPO / "reports/partner_data_recovery/outputs/_final_closeout"
TEMPLATE = REPO / "reports/partner_data_recovery/inputs/Round 1_run/2026-10-01_master_log_main_edited.xlsx"
FRAME_DIR = REPO / "input_data/sampling_frame"
# Round 1's frame (archived 9 Oct) - fallback only, for a cluster no longer in the current frame
FRAME_V14 = FRAME_DIR / "_archive_2026-10-09_pre_v17_sync/NGA_MSNA_2026_stage2_sampling_frame_v14_FULL.csv"
IDENT_COLS = ["uuid", "submission_date", "enum_id", "admin1", "admin2", "org_id", "pop_type", "hh_size", "strata_id", "cluster_id",
              "sample_point_id", "households_in_cluster", "households_in_cluster_estimate", "target_households", "reserve_households",
              "total_to_draw_planned"]
R1_LOG = CL / "MSNA_N-WEC_2026_Round1_deletion_log_2026-10-02.csv"
# identity columns describe the data officer's dataset as HE has it - before any of this log's corrections
PRE_CORRECTION_SUBS = REPO / "data/_archive/2026-10-02_pre_round1_corrections_overlay/real_submissions.csv"
TODAY = datetime.date.today().isoformat()
DO_COLS = ["uuid", "submission_date", "enum_id", "admin1", "admin2", "org_id", "pop_type", "hh_size", "strata_id", "cluster_id",
           "sample_point_id", "households_in_cluster", "households_in_cluster_estimate", "target_households", "reserve_households",
           "total_to_draw_planned", "check_binding", "check_id", "issue", "label", "question", "old_value", "change_type",
           "new_value", "Data_Feedback", "FO_Comments", "status", "response_file", "check_date"]
EXTRA = ["deletion_round", "interview_round", "decision_basis", "decided_by", "tracker_issue_id", "review_reminder",
         "already_in_do_master_log"]
SUFFIX = re.compile(r"_[b-z]$")
ORG = {"nrc": "NRC", "coopi": "COOPI", "imc": "IMC", "plan": "PLAN", "malteser": "Malteser", "fact": "FACT", "mdm": "MdM",
       "acf": "ACF", "crs": "CRS", "care": "CARE", "si": "SI", "lhi": "LHI", "intersos": "INTERSOS", "zoa": "ZOA",
       "street_child": "Street Child"}
CLOSEOUT_R1 = "MSNA team - Round 1 closeout (2026-10-02)"
CLOSEOUT = f"MSNA team - final closeout ({TODAY})"
CLOSEOUT_PREFIX = re.compile(r"^(?:Round 1|Final) closeout(?: \([^)]*\))?:\s*")
DURATION_RULE_MS = 1_197_000  # 19.95 minutes: the lowest audit-trail duration that rounds to 20.0 at one decimal place


def latest_frame():
    vs = sorted(FRAME_DIR.glob("NGA_MSNA_2026_stage2_sampling_frame_v*_FULL.csv"),
                key=lambda p: int(re.search(r"_v(\d+)_FULL", p.name).group(1)))
    if not vs:
        raise SystemExit(f"no FULL stage-2 frame in {FRAME_DIR}")
    return vs[-1]


def rule_label(res, rule):
    """'Round 1 closeout rule Q3-A/B (Jack)' for a Round 1 decision; the same rule applied to the final dataset (16 Oct,
    Jack #67) says so."""
    if (res or "").startswith("Final closeout"):
        return f"Closeout rule {rule} (Jack, 2 Oct 2026; applied to the final dataset, 16 Oct 2026)"
    return f"Round 1 closeout rule {rule} (Jack)"


def miss(x):
    return x is None or str(x).strip() in ("", "NA")


def clean(x):
    s = "" if x is None else str(x).strip()
    return "" if s in ("NA", "None") or s.startswith("NA_") else (s[:-2] if re.fullmatch(r"\d+\.0", s) else s)


def org(o):
    return ORG.get(o, (o or "").upper().replace("_", " "))


def fmt_min(ms):
    """Audit-trail minutes as the master log writes them (1 decimal) - but never let rounding print a value that is
    under 20 minutes as '20.0' (three such interviews are 19.9x minutes)."""
    v = ms / 60000
    for nd in (1, 2, 3):
        s = f"{v:.{nd}f}"
        if not (float(s) >= 20 > v):
            return s
    return f"{v:.4f}"


def claim_field(r):
    if clean(r.get("pop_type")) == "idp":
        if not miss(r.get("idp_hh_number_from_listing")):
            return "idp_hh_number_from_listing", clean(r.get("idp_hh_number_from_listing"))
        if not miss(r.get("idp_walk_position")):
            return "idp_walk_position", clean(r.get("idp_walk_position"))
    return "non_idp_point_id", clean(r.get("non_idp_point_id"))


def strip_prefix(res):
    return cap(CLOSEOUT_PREFIX.sub("", res).strip())


def cap(s):
    s = s.strip()
    return s[:1].upper() + s[1:]


def partner_answer(t):
    """Early partner answers (Sep 2026, before the tracker recorded confirmed_by) hold the partner's own value as the
    whole resolution - a household number or a point id."""
    return t["confirmed_by"] == "partner" or (t["issue_type"] in ("gps_duplicate", "idp_listing_duplicate") and miss(t["confirmed_by"]))


def source_for(t):
    res = t["resolution"] or ""
    if res.startswith("Final closeout"):
        return CLOSEOUT
    if res.startswith("Round 1 closeout") or not miss(t["fallback_status"]):
        return "Partner recovery workbook (Round 1, returned Sep/Oct 2026)" if t["confirmed_by"] == "partner" else CLOSEOUT_R1
    if "independently computed" in res:
        return "MSNA team - independent check (audit-trail duration / consent answer)"
    if "legacy exclusions file" in res:
        return "MSNA team - quality exclusions list of 30 Aug 2026"
    if partner_answer(t):
        return "Partner recovery workbook (earlier round)"
    return f"MSNA team - recovery tracker ({t['resolution_date'] or 'earlier round'})"


def basis_for_kept(t):
    mech = t["fallback_mechanism"] or ""
    if t["deletion_reason"] == "fcs_zero":
        return "Policy decision (Jack, 10 Sep 2026): fcs_zero is a logical-error flag, not a removal reason"
    if mech == "live_claimant_rule":
        return rule_label(t["resolution"], "Q3-D") + ": no other live claimant left -> not a duplicate"
    if mech == "keep_distinct_household_unchanged":
        return rule_label(t["resolution"], "Q3-C") + ": different household -> kept; identifier left unchanged"
    if mech == "q4_listing_check":
        return "Round 1 closeout rule Q4 (Jack): household number checked against the real listing"
    if mech == "q6_date_exception" or t["deletion_reason"] == "date_outlier":
        return "Round 1 closeout rule Q6 (Jack): wrong device date -> kept as an exception"
    if mech == "duration_rule_one_decimal":
        return "Duration rule (Jack, 2 Oct 2026): removed only when under 20.0 minutes at one decimal place (under 19.95 minutes)"
    if partner_answer(t):
        return "Partner response (recovery workbook)"
    if "live claimant" in (t["resolution"] or "").lower() or "live_claimant" in (t["resolution"] or ""):
        return "Live-claimant rule (Jack/Coordinator, 25 Sep 2026): the first live claimant of a point is not a duplicate"
    return "MSNA team review (recovery tracker)"


def main():
    preview = "--preview" in sys.argv
    out_dir = Path(sys.argv[sys.argv.index("--out") + 1]) if "--out" in sys.argv else OUT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    r1members = {r["submission_uuid"] for r in csv.DictReader(open(REPO / "data/ROUND1_MEMBERSHIP.csv", encoding="utf-8-sig"))}
    pre = {r["submission_uuid"]: r for r in csv.DictReader(open(PRE_CORRECTION_SUBS, encoding="utf-8-sig"))}
    post = {r["submission_uuid"]: r for r in csv.DictReader(open(REPO / "data/real_submissions.csv", encoding="utf-8-sig"))}
    members = set(post)  # the WHOLE final dataset
    missing_r1 = r1members - members
    if missing_r1:
        raise SystemExit(f"{len(missing_r1)} Round 1 submission(s) are not in the final data - check the export before building")
    trk = list(csv.DictReader(open(REPO / "reports/partner_data_recovery/scripts/recovery_issue_tracker.csv", encoding="utf-8-sig")))
    by_id = {t["issue_id"]: t for t in trk}
    # settled removals straight from the tracker (exactly build_confirmed_deletions_overlay.R's rule), cross-checked
    # against the overlay file itself so a stale overlay can never slip a different set into the log
    removed = {t["uuid"] for t in trk if t["issue_type"] == "confirmed_deletion" and t["status"] in ("confirmed", "contested")
               and miss(t["recovery_type"])} & members
    overlay = {r["uuid"] for r in csv.DictReader(open(REPO / "data/CONFIRMED_DELETIONS_OVERLAY.csv", encoding="utf-8-sig"))
               if r["status"] in ("confirmed", "contested")} & members
    if overlay != removed:
        raise SystemExit(f"CONFIRMED overlay is stale vs the tracker ({len(overlay)} vs {len(removed)}) - rebuild it first")
    audit = {r["uuid"]: float(r["duration_audit_sum_all_ms"])
             for r in csv.DictReader(open(REPO / "cleaning/real/audit_duration_cache.csv", encoding="utf-8-sig"))
             if not miss(r["duration_audit_sum_all_ms"])}
    nrc_contested = {r["key"] for r in csv.DictReader(open(CL / "round1_returned_workbooks_classified.csv", encoding="utf-8"))
                     if r["org_id"] == "nrc" and r["sheet"] == "Confirmed Deletions" and r["deletion_reason"] == "duration_under_20"
                     and r["partner_answer"].strip().lower() == "yes"}
    corr = list(csv.DictReader(open(CL / "round1_final_corrections.csv", encoding="utf-8")))
    clusters, strata_of, idp_cat = {}, {}, {}
    frame = latest_frame()
    for path in (frame, FRAME_V14):  # current frame first; Round 1's v14 only fills clusters the current one lacks
        if not path.exists():
            continue
        for r in csv.DictReader(open(path, encoding="utf-8-sig")):
            if r["cluster_id"] in strata_of and path != frame:
                continue
            clusters.setdefault(r["cluster_id"], r)
            strata_of.setdefault(r["cluster_id"], r["strata_id"])
            if r["pop_type"] == "idp" and not miss(r["idp_population_category"]):
                idp_cat.setdefault(r["cluster_id"], r["idp_population_category"])
    # what the Round 1 log (as sent, 2 Oct) removed - deletion_round, and the reversal check below - and its identity
    # columns, which a Round 1 interview keeps verbatim here (its frame values can't drift with later frame versions)
    r1_log = list(csv.DictReader(open(R1_LOG, encoding="utf-8-sig")))
    r1_removed = {r["uuid"] for r in r1_log if r["change_type"] == "remove_survey"}
    r1_ident = {}
    for r in r1_log:
        r1_ident.setdefault(r["uuid"], {c: r[c] for c in IDENT_COLS})
    exceptions = []

    def exc(kind, blocking, u, detail, issue_id=""):
        exceptions.append(dict(kind=kind, blocking="yes" if blocking else "no", uuid=u, org_id=clean(post.get(u, {}).get("org_id")),
                               interview_round="Round 1" if u in r1members else "Round 2", tracker_issue_id=issue_id, detail=detail))

    # the master log, to mark what it already does for the same interviews
    wb = openpyxl.load_workbook(TEMPLATE, read_only=True, data_only=True)
    ws = wb["cleaning_log"]
    ws.reset_dimensions()
    rows = list(ws.iter_rows(values_only=True))
    ix = {h: i for i, h in enumerate(rows[0])}
    body = [list(r) + [None] * (len(rows[0]) - len(r)) for r in rows[1:]]
    do_removed = {r[ix["uuid"]]: r[ix["check_id"]] for r in body if r[ix["change_type"]] == "remove_survey"}
    do_changed = {(r[ix["uuid"]], r[ix["question"]]): (clean(r[ix["new_value"]]), r[ix["check_id"]])
                  for r in body if r[ix["change_type"]] == "change_response"}
    do_dur = {r[ix["uuid"]]: r[ix["old_value"]] for r in body if r[ix["check_id"]] == "duration_low"}
    do_dates = {r[ix["uuid"]] for r in body if r[ix["check_id"]] == "device_date_error"}
    sheet_rows = {}
    for name in ("readme", "validation_rules"):
        s = wb[name]
        s.reset_dimensions()
        sheet_rows[name] = [list(r) for r in s.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]

    def ident(u):
        # Round 1 interviews: verbatim from the Round 1 log where they're in it, else as it would have built them (before
        # its corrections); later interviews: the final data (no corrections apply to them)
        if u in r1_ident:
            return dict(r1_ident[u])
        r = pre.get(u) if (u in r1members and u in pre) else post.get(u, {})
        pop = clean(r.get("pop_type"))
        cl = clean(r.get("matched_cluster_id"))
        c = clusters.get(cl, {})
        idp = pop == "idp"
        hic = c.get("households_in_cluster", "")
        tgt, rsv = (c.get("target_households", ""), c.get("reserve_households", "")) if idp else ("", "")
        total = str(int(float(tgt)) + int(float(rsv))) if idp and not miss(tgt) and not miss(rsv) else ""
        return dict(uuid=u, submission_date=clean(r.get("submission_date")), enum_id=clean(r.get("enum_id")),
                    admin1=clean(r.get("admin1")), admin2=clean(r.get("admin2_submitted")), org_id=clean(r.get("org_id")),
                    pop_type=pop, hh_size=clean(r.get("hh_size")), strata_id=clean(r.get("matched_strata_id")), cluster_id=cl,
                    sample_point_id=cl if idp else clean(r.get("non_idp_point_id")), households_in_cluster=hic,
                    households_in_cluster_estimate=hic if idp else "", target_households=tgt, reserve_households=rsv,
                    total_to_draw_planned=total)

    out = []
    referenced = set()

    def add(row, issue_ids):
        base = {c: "" for c in DO_COLS + EXTRA}
        base.update(row)
        base["status"] = "addressed"
        base["tracker_issue_id"] = ";".join(i for i in issue_ids if i)
        base["interview_round"] = "Round 1" if base["uuid"] in r1members else "Round 2"
        if base["change_type"] == "remove_survey":
            base["deletion_round"] = "Round 1" if base["uuid"] in r1_removed else "Round 2"
        referenced.update(i for i in issue_ids if i)
        out.append(base)

    # ================= remove_survey ===============================================================================
    cd_row = {t["uuid"]: t for t in trk if t["issue_type"] == "confirmed_deletion" and t["uuid"] in removed}
    assert set(cd_row) == removed, "a removed interview has no confirmed_deletion tracker row"
    for u in sorted(removed):
        t = cd_row[u]
        r = pre.get(u) if (u in r1members and u in pre) else post.get(u, {})
        reason, res, mech = t["deletion_reason"], t["resolution"] or "", t["fallback_mechanism"] or ""
        ms = audit.get(u)
        # THE DURATION RULE (Jack, 2 Oct 2026): removed only when the audit-trail duration, rounded to one decimal place
        # of a minute, is below 20.0 - i.e. under 19.95 minutes (1,197,000 ms), the master log's own rounding
        short = ms is not None and ms < DURATION_RULE_MS
        note, reminder = "", ""
        if reason == "no_consent":
            cid, issue, q, old = "no_consent", "Respondent did not consent", "consent", "no"
            feedback = "Consent recorded as no."
            fo = "Survey removed. The respondent did not give consent to be interviewed, so none of its answers can be used."
            basis = "Mandatory removal: no consent (automatic, no appeal - Jack, 11 Sep 2026)"
        # the data officer's priority order: after consent, ANY interview under 19.95 minutes is a duration removal
        # (since 4/12 Oct the duration rule overrides a recovery or an open flag of another reason - its notes say so)
        elif reason == "duration_under_20" or (t["status"] == "contested" and miss(reason)) or short:
            cid, issue, q = "duration_under_20", "Interview shorter than 20 minutes", "duration_audit_sum_all_minutes"
            old = fmt_min(ms) if ms is not None else ""
            feedback = (f"Interview lasted {old} minutes by its audit trail (time actually spent on the questions), under the "
                        f"20-minute minimum.")
            fo = (f"Survey removed. At {old} minutes the questionnaire cannot have been administered in full; interviews under "
                  f"20 minutes are removed automatically, with no appeal.")
            basis = "Validated methodology threshold: under 20 minutes by the audit trail = removed, no appeal (Jack, 6/11 Sep 2026)"
            if u in nrc_contested:
                note = (" NRC contested this removal in its Round 1 recovery workbook; not accepted (Jack, 1 Oct 2026: every "
                        "interview under 20 minutes is deleted).")
            if t["status"] == "contested":
                why = re.sub(r"^.*Partner reason:\s*", "", res).strip()
                note = (f" {org(t['org_id'])} contested the removal (\"{why[:150]}{'...' if len(why) > 150 else ''}\"); the "
                        f"contest was reviewed and rejected on {t['resolution_date']}, so the removal stands.")
            if reason == "duplicate_point" or "duplicate_point" in (t["notes"] or ""):
                note = (" It had also been flagged for sharing its sample point with another interview; the removal rests "
                        "on its duration alone.")
            elif not miss(reason) and reason != "duration_under_20":
                note = f" It had also been flagged for {reason.replace('_', ' ')}; the removal rests on its duration alone."
            if not short:
                exc("duration_removal_not_supported", True, u,
                    f"removed for duration but the audit trail gives {'no duration' if ms is None else fmt_min(ms) + ' minutes'} - "
                    f"the duration rule removes only under 19.95 minutes", t["issue_id"])
        elif mech == "q3_same_household":
            cid, issue = "duplicate_point_same_household", "Same household interviewed twice"
            q, old = claim_field(r)
            feedback = (f"An earlier interview recorded at the same {'sample point' if q == 'non_idp_point_id' else 'household number'} "
                        f"({old}) has the same household size and head of household (sex, and age within 2 years).")
            fo = "Survey removed. " + strip_prefix(res)
            basis = rule_label(res, "Q3-A/B") + ": same household interviewed twice -> keep the first, remove the later"
        elif mech == "not_at_sampled_location":
            cid, issue, q, old = "not_at_sampled_location", "Interview not at a sampled location", "non_idp_point_id", clean(r.get("non_idp_point_id"))
            feedback = "Device GPS far from every drawn point in the interview's stratum."
            fo = "Survey removed. " + cap(re.sub(r"^Round 1 closeout \(Jack, 2 Oct\):\s*", "", res)).replace("; deleted.", ".")
            basis = "Round 1 closeout rule (Jack, 2 Oct): more than 2 km from every drawn point -> removed"
        elif reason == "duplicate_point":
            cid, issue = "duplicate_point", "Second interview at an occupied sample point"
            q, old = claim_field(r)
            feedback = f"Another completed interview had already been recorded at {old}."
            fo = (f"Survey removed. A second completed interview at the same sample point as an earlier one; {org(t['org_id'])} "
                  f"did not contest the removal (recovery workbook, {t['resolution_date']}).")
            basis = "Partner response: removal not contested"
        elif reason == "gps_no_match_partner_confirmed":
            cid, issue, q, old = "gps_no_match", "No household match at the sample point", "non_idp_point_id", clean(r.get("non_idp_point_id"))
            feedback = "Device GPS far from the claimed sample point; no nearby drawn household matched."
            fo = "Survey removed. " + res.replace("ACF GPS Duplicates response: ", "ACF's GPS Duplicates response: ") + "."
            basis = "Partner response: removal confirmed by the partner"
        elif reason == "crs_unmatched":
            cid, issue, q, old = "sampling_assignment_missing", "Sample point, cluster and population group not recorded", "sample_point_id", ""
            feedback = "The sampling section was not completed on the device, and its GPS cannot be placed in the sample."
            fo = ("Survey removed. The sample point, cluster and population group were never recorded, and the " +
                  re.sub(r"^Round 1 closeout \(Jack, Q5\):\s*device", "device", res).replace(" - cannot be placed; deleted.",
                                                                                              ", so it cannot be placed in the sample."))
            basis = "Round 1 closeout rule Q5 (Jack): place by device GPS; nothing within 150 m -> removed"
        else:
            raise SystemExit(f"no rule to describe the removal of {u}: {reason} / {t['status']} / {res[:80]}")
        if u in do_removed:
            in_do = f"yes - the master log also removes this survey ({do_removed[u]})"
        elif u in do_dur and cid == "duration_under_20":
            in_do = (f"CONFLICT - the master log keeps it as duration_low ({do_dur[u]} minutes); this log removes it: the "
                     f"MSNA team's audit-trail duration is {old} minutes")
        elif u in do_dur:
            in_do = (f"CONFLICT - the master log keeps it (duration_low, {do_dur[u]} minutes, no action); this log removes it "
                     f"for a different reason: {issue.lower()}")
        else:
            in_do = "no"
        add(dict(ident(u), check_binding=f"{cid} ~/~ {u}", check_id=cid, issue=issue, label=q, question=q, old_value=old,
                 change_type="remove_survey", new_value="", Data_Feedback=feedback, FO_Comments=fo + note,
                 response_file=source_for(t), check_date=t["detected_date"], decision_basis=basis,
                 decided_by=t["confirmed_by"] or "internal_team", review_reminder=reminder, already_in_do_master_log=in_do),
            [t["issue_id"]])

    # ================= change_response =============================================================================
    by_u = defaultdict(dict)
    for c in corr:
        if c["uuid"] in members and c["uuid"] not in removed:
            by_u[c["uuid"]][c["field"]] = c
    expand = {"non_idp_point_id": ["non_idp_point_id", "sample_point_id"], "idp_hh_number_from_listing": ["idp_hh_number_from_listing"],
              "idp_walk_position": ["idp_walk_position"], "pop_type": ["sample_pop_type_filter", "pop_type"],
              "sample_point_id": ["sample_point_id"], "dist_btn_sample_collected": ["dist_btn_sample_collected"]}
    for u, fields in sorted(by_u.items()):
        placed = "pop_type" in fields  # GPS placement of an interview whose sampling section was never filled (Q5)
        for f, c in fields.items():
            if f == "dist_btn_sample_collected" and clean(c["old_value"]) == c["new_value"]:
                continue  # the new point sits exactly where the old one was (co-located drawn points): distance unchanged
            ids = [i for i in c["issue_ids"].split(";") if i]
            vars_ = ["cluster_id"] + (["idp_cluster_id"] if c["new_value"].startswith("idp_") else []) if f == "cluster_id" else expand[f]
            for v in vars_:
                add(change_row(u, v, c["old_value"], c["new_value"], c, ids, ident, by_id, do_changed, do_removed), ids)
        if placed:
            c0 = fields["cluster_id"]
            cl, pop = c0["new_value"], fields["pop_type"]["new_value"]
            tg = "non_idp" if pop == "non_idp" else idp_cat.get(cl, "")
            derived = [("strata_id", strata_of.get(cl, "")), ("target_group", tg)] + ([("idp_category", tg)] if pop == "idp" else [])
            ids = [i for i in c0["issue_ids"].split(";") if i]
            for v, new in derived:
                cc = dict(c0, field=v, justification=f"Set with the GPS placement at {cl} (the value the form derives from it). " + c0["justification"])
                add(change_row(u, v, "", new, cc, ids, ident, by_id, do_changed, do_removed), ids)

    # ================= no_action ===================================================================================
    for t in trk:
        u = t["uuid"]
        if miss(u) or u not in members or u in removed or t["issue_id"] in referenced:
            continue
        if t["status"] not in ("confirmed", "contested"):
            exc("open_tracker_row", True, u, f"{t['issue_type']} / {t['deletion_reason'] or '-'} still {t['status']} - "
                f"close it (final closeout) before the log is final", t["issue_id"])
            referenced.add(t["issue_id"])
            continue
        r = pre.get(u) if (u in r1members and u in pre) else post.get(u, {})
        reason = t["deletion_reason"] or t["issue_type"]
        res = re.sub(r"\s*REVIEW:.*$", "", t["resolution"] or "")
        q, old, in_do = "", "", "no"
        if reason == "duplicate_point" and (t["fallback_mechanism"] or "") == "keep_distinct_household_unchanged":
            q, old = claim_field(r)
            issue = "Sample point / household number shared with a different household"
            feedback = ("Another interview, of a different household (household size or head's sex/age differ), holds the same "
                        "identifier.")
            fo = "No action. " + strip_prefix(res)
        elif reason == "duplicate_point":
            q, old = claim_field(r)
            issue, feedback = "Sample point shared with another interview", f"Another interview was recorded at {old}."
            fo = "No action. " + strip_prefix(res)
        elif reason == "listing_missing":
            q, old = "idp_hh_number_from_listing", clean(r.get("idp_hh_number_from_listing"))
            issue = "Household number not checked against a household listing"
            feedback = "The cluster had no household listing when the interview was flagged."
            fo = "No action. " + strip_prefix(res)
        elif reason == "date_outlier":
            q, old = "start", clean(r.get("start_datetime"))
            issue, feedback = "Implausible interview date", "The device clock gave an impossible interview date."
            fo = "No action. " + strip_prefix(res) + " The responses are retained."
            if u in do_dates:
                in_do = "yes - the master log corrects today/start/end for this interview (device_date_error)"
        elif reason == "duration_under_20":
            # reinstated under the duration rule (Jack, 2 Oct 2026) - the master log keeps these too (duration_low)
            q, old = "duration_audit_sum_all_minutes", (f"{audit[u] / 60000:.2f}" if u in audit else "")
            issue = "Interview flagged under 20 minutes - not short under the duration rule"
            feedback = f"Audit-trail duration {old} minutes (time actually spent on the questions)."
            fo = "No action - the interview is retained. " + strip_prefix(res)
            if u in do_dur:
                in_do = f"yes - the master log keeps it too (duration_low, {do_dur[u]} minutes)"
        elif reason == "fcs_zero":
            issue, feedback = "Food consumption score of zero", "Every food group recorded as eaten on zero days."
            fo = ("No action - no longer a removal reason (Jack, 10 Sep 2026): the interview is retained; the FCS module "
                  "itself is handled by the master log's own flag_fcs_zero check.")
        elif t["issue_type"] in ("gps_duplicate", "idp_listing_duplicate"):
            q, old = claim_field(r)
            issue = ("Sample point shared with another interview" if t["issue_type"] == "idp_listing_duplicate"
                     else "Device GPS far from the claimed point")
            feedback = issue + "."
            fo = (f"No action. {org(t['org_id'])} confirmed the recorded "
                  f"{'household number' if q != 'non_idp_point_id' else 'sample point'} ({old}) is correct.")
        else:
            raise SystemExit(f"no rule to describe the retained flag {t['issue_id']} ({reason})")
        add(dict(ident(u), check_binding=f"{reason} ~/~ {u}", check_id=reason, issue=issue, label=q, question=q, old_value=old,
                 change_type="no_action", new_value="", Data_Feedback=feedback, FO_Comments=fo, response_file=source_for(t),
                 check_date=t["detected_date"], decision_basis=basis_for_kept(t),
                 decided_by="partner" if partner_answer(t) else (t["confirmed_by"] or "internal_team"),
                 already_in_do_master_log=in_do), [t["issue_id"]])
    dtm = {t["cluster_id"]: t for t in trk if t["issue_type"] == "missing_hh_listing" and t["fallback_mechanism"] == "iom_dtm_listing"}
    for u in sorted(members - removed):
        r = post[u]
        t = dtm.get(r["matched_cluster_id"])
        if t is None or r["interview_outcome"] != "completed":
            continue
        site = re.search(r"IOM DTM site (\S+?)\.? used", t["resolution"] or "")
        add(dict(ident(u), check_binding=f"missing_hh_listing ~/~ {u}", check_id="missing_hh_listing",
                 issue="No household listing was ever submitted for this cluster", label="sample_point_id", question="sample_point_id",
                 old_value=r["matched_cluster_id"], change_type="no_action", new_value="",
                 Data_Feedback=f"Cluster {r['matched_cluster_id']} has no Household Listing submission.",
                 FO_Comments=(f"No action. No household listing was submitted for cluster {r['matched_cluster_id']}; the IOM DTM "
                              f"site {site.group(1) if site else ''} is used as its listing reference. This concerns the listing "
                              f"record for the cluster rather than this interview, whose responses are retained."),
                 response_file=CLOSEOUT if (t["resolution"] or "").startswith("Final closeout") else CLOSEOUT_R1,
                 check_date=t["detected_date"],
                 decision_basis=rule_label(t["resolution"], "Q8") + ": no listing -> IOM DTM site as the listing reference",
                 decided_by="internal_team", already_in_do_master_log="no"), [t["issue_id"]])
    open_cluster = [t for t in trk if t["issue_type"] == "missing_hh_listing" and t["status"] not in ("confirmed", "contested")]
    for t in open_cluster:
        exc("open_tracker_row", True, "", f"missing_hh_listing for cluster {t['cluster_id']} still {t['status']} - close it "
            f"(rule Q8) before the log is final", t["issue_id"])

    # ================= checks ======================================================================================
    unref = [t["issue_id"] for t in trk if not miss(t["uuid"]) and t["uuid"] in members and t["uuid"] not in removed
             and t["issue_id"] not in referenced]
    assert not unref, f"{len(unref)} tracker row(s) not represented in the log: {unref[:5]}"
    # (1) a Round 1 removal that the final data no longer removes - reported, never dropped silently
    for u in sorted(r1_removed - removed):
        ts = [t for t in trk if t["uuid"] == u and t["issue_type"] == "confirmed_deletion"]
        exc("round1_removal_reversed", False, u, "removed in the Round 1 log, retained now: " +
            ("; ".join(f"{t['status']}/{t['recovery_type'] or '-'}: {(t['resolution'] or '')[:160]}" for t in ts) or "no tracker row"),
            ";".join(t["issue_id"] for t in ts))
    # (2) #68 (Jack, 16 Oct): every Achieved interview needs its audit-trail duration - no placeholder. An audit trail
    # cleaningtools can't total (a question with no end/start time) passes when its MINIMUM - the complete rows alone,
    # final_audit_consistency.py - already reaches 19.95 minutes (Jack's decision, 16 Oct): a kept row states it.
    lb_files = sorted(OUT_DIR.glob("final_audit_incomplete_lower_bounds_*.csv"))
    lower = {r["uuid"]: r for r in csv.DictReader(open(lb_files[-1], encoding="utf-8-sig"))} if lb_files else {}
    for u in sorted(members - removed):
        if post[u]["interview_outcome"] != "completed" or u in audit:
            continue
        b = lower.get(u)
        if b and float(b["lower_bound_ms"]) >= DURATION_RULE_MS and int(b["negative_rows"]) == 0:
            mn = f"{float(b['lower_bound_ms']) / 60000:.1f}"
            n = int(b["incomplete_rows"])
            add(dict(ident(u), check_binding=f"duration_audit_incomplete ~/~ {u}", check_id="duration_audit_incomplete",
                     issue="Audit trail incomplete - interview duration cannot be totalled",
                     label="duration_audit_sum_all_minutes", question="duration_audit_sum_all_minutes", old_value="",
                     change_type="no_action", new_value="",
                     Data_Feedback=(f"{n} question{'s' if n > 1 else ''} in the audit trail ha{'ve' if n > 1 else 's'} no end "
                                    f"or start time, so the audit-trail duration cannot be totalled; the questions with "
                                    f"both times recorded add up to {mn} minutes."),
                     FO_Comments=(f"No action - the interview is retained. Its audit-trail duration is at least {mn} minutes "
                                  f"(the missing times can only add to it), so it is not short under the duration rule."),
                     response_file=CLOSEOUT, check_date=TODAY,
                     decision_basis=("Duration rule (Jack, 2 Oct 2026) on an incomplete audit trail: passes when the recorded "
                                     "questions alone reach 19.95 minutes (Jack, 16 Oct 2026)"),
                     decided_by="internal_team", already_in_do_master_log="no"), [])
        else:
            exc("no_audit_duration", True, u, "Achieved interview with no audit-trail duration in the audit export" +
                (f" (incomplete audit trail, minimum {float(b['lower_bound_ms']) / 60000:.2f} minutes)" if b else "") +
                " - the duration rule cannot be applied; Jack: wait for the full audit")
    for o in out:
        assert o["change_type"] in ("change_response", "blank_response", "remove_survey", "no_action") and o["uuid"] in members, o
        for k in ("check_binding", "check_id", "issue", "Data_Feedback", "FO_Comments", "response_file", "decision_basis", "decided_by"):
            assert str(o[k]).strip(), f"empty {k}: {o['uuid']} {o['check_id']} {o['question']}"
        if o["change_type"] == "change_response":
            assert o["new_value"] not in ("", None) and o["new_value"] != o["old_value"], f"no-op change: {o['uuid']} {o['question']}"
    keys = Counter((o["uuid"], o["question"]) for o in out if o["change_type"] == "change_response")
    dup = [k for k, n in keys.items() if n > 1]
    assert not dup, f"two change rows for one variable: {dup[:5]}"
    binds = Counter(o["check_binding"] for o in out)
    assert max(binds.values()) == 1, f"check_binding not unique: {[k for k, n in binds.items() if n > 1][:5]}"
    assert {o["uuid"] for o in out if o["change_type"] == "remove_survey"} == removed
    assert not ({o["uuid"] for o in out if o["change_type"] == "change_response"} & removed)
    order = {"remove_survey": 0, "change_response": 1, "no_action": 2}
    out.sort(key=lambda o: (o["submission_date"] or "9999", o["uuid"], order[o["change_type"]], o["question"]))
    stem = out_dir / f"MSNA_N-WEC_2026_FINAL_deletion_log_{TODAY}"
    with open(f"{stem}_EXCEPTIONS.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["kind", "blocking", "uuid", "org_id", "interview_round", "tracker_issue_id", "detail"])
        w.writeheader()
        w.writerows(exceptions)
    print(f"exceptions: {len(exceptions)} -> {stem.name}_EXCEPTIONS.csv")
    for (k, b), n in sorted(Counter((e["kind"], e["blocking"]) for e in exceptions).items()):
        print(f"  {n:5d}  {k} (blocking: {b})  by org: {dict(Counter(e['org_id'] for e in exceptions if e['kind'] == k).most_common())}")
    r1_log_cid = {r["uuid"]: r["check_id"] for r in csv.DictReader(open(R1_LOG, encoding="utf-8-sig")) if r["change_type"] == "remove_survey"}
    relabel = Counter((r1_log_cid[o["uuid"]], o["check_id"]) for o in out
                      if o["change_type"] == "remove_survey" and o["uuid"] in r1_log_cid and r1_log_cid[o["uuid"]] != o["check_id"])
    print(f"Round 1 removals labelled differently from the Round 1 log (same removal, other basis): {dict(relabel)}")
    blocking = [e for e in exceptions if e["blocking"] == "yes"]
    if blocking and not preview:
        raise SystemExit(f"{len(blocking)} blocking exception(s) - see {stem.name}_EXCEPTIONS.csv; nothing else written")
    write(out, sheet_rows, stem, preview, r1members)


def change_row(u, var, old, new, c, ids, ident, by_id, do_changed, do_removed):
    t = by_id.get(ids[0], {}) if ids else {}
    reason = (t.get("deletion_reason") or t.get("issue_type") or "") if t else ""
    dec = c["decision"]
    if c["field"] == "dist_btn_sample_collected":
        cid, issue = "sample_point_changed", "GPS distance to follow the new sample point"
        feedback = "The interview's sample point changed (see its sample_point_id row); the recorded distance belonged to the old point."
        basis = "Follows the sample point change (the form's own distance formula)"
    elif reason == "crs_unmatched" or "PLACE" in dec:
        cid, issue = "sampling_assignment_missing", "Sample point, cluster and population group not recorded"
        feedback = "The sampling section was not completed on the device; the interview was placed in the sample by its device GPS."
        basis = "Round 1 closeout rule Q5 (Jack): place by device GPS (within 150 m of a drawn point or IDP site)"
    elif dec.startswith("final uniqueness"):
        cid, issue = "duplicate_key_uniqueness", "Two different households share one sample point / household number"
        feedback = "After all other corrections, two interviews of different households still shared this identifier."
        basis = "Round 1 closeout rule Q3 (Jack): distinct households both kept, the later one suffixed"
    elif c["decided_by"] == "partner" and c["sources"] == "step2":
        cid = reason or "duplicate_point"
        issue = "Sample point / household number corrected by the partner"
        feedback = "The recorded sample point / household number was shared with another interview."
        basis = "Partner response, Round 1 recovery workbook (checked against the data before acceptance)"
    elif c["sources"].startswith("earlier"):
        cid = reason or "partner_correction_earlier_round"
        issue = "Sample point / household number corrected by the partner"
        feedback = "The recorded sample point / household number was flagged in an earlier recovery round."
        basis = ("Partner response, earlier recovery round (recorded in the MSNA tracker, not yet applied)" if c["sources"] == "earlier"
                 else "Partner response, earlier recovery round; then Round 1 closeout rule Q3-C (Jack): that number is held by a "
                      "different household -> nearest free listed household")
    elif reason == "listing_missing":
        cid, issue = "listing_missing", "Household number not in the cluster's household listing"
        feedback = "The recorded household number is not among the drawn households of the cluster's household listing."
        basis = "Round 1 closeout rule Q4 (Jack): number not in the real listing -> nearest free listed number"
    elif reason == "gps_duplicate":
        cid, issue = "gps_duplicate", "Device GPS far from the claimed sample point"
        feedback = "The device GPS is far from the claimed point, and a different household already held it."
        basis = "Round 1 closeout rule Q3 (Jack): recorded at the drawn point nearest the device, suffixed"
    elif "REASSIGN" in dec:
        cid, issue = "duplicate_point", "Sample point / household number shared with a different household"
        feedback = "Another interview, of a different household (household size or head's sex/age differ), holds the same identifier."
        basis = "Round 1 closeout rule Q3-C (Jack): different household -> nearest free household"
    elif "KEEP" in dec:
        cid, issue = "duplicate_point", "Sample point / household number shared with a different household"
        feedback = ("Another interview, of a different household (household size or head's sex/age differ), holds the same "
                    "identifier, and no free household was left within the limits.")
        basis = "Round 1 closeout rule Q3-C (Jack): different household, no free household -> both kept, the later one suffixed"
    else:
        raise SystemExit(f"no rule to describe the change {u} {var} ({dec})")
    # old_value = the value in the data officer's dataset at the point this row applies. Where his own log recovers a
    # sampling field first (auto_recovery_sampling_fields - the NG037 form error), this row FOLLOWS that recovery: its
    # old value is his recovered value, and non_idp_point_id (which his recovery doesn't set) is still blank there.
    eff_old = clean(old)
    auto_recovered = do_changed.get((u, "sample_point_id"), ("", ""))[1] == "auto_recovery_sampling_fields"
    if var == "non_idp_point_id" and auto_recovered:
        eff_old = ""
    his = do_changed.get((u, var))
    if u in do_removed:
        in_do = (f"CONFLICT - the master log removes this survey ({do_removed[u]}); this log keeps it, placed by its device "
                 f"GPS: drop that removal and apply these changes")
    elif his and his[0] == new:
        in_do = "yes - same change"
    elif his and eff_old in ("", his[0]):
        in_do = f"follows - the master log first sets {var} to {his[0]} ({his[1]}); apply this change after it"
        eff_old = his[0]
    elif his:
        in_do = f"CONFLICT - the master log sets {var} to {his[0]}; this log's value supersedes it"
    else:
        in_do = "no"
    # every correction is a Round 1 decision (2 Oct) - the final round changed no identifiers
    src_part = {"closeout": CLOSEOUT_R1, "unique": CLOSEOUT_R1, "step2": "Partner recovery workbook (Round 1)",
                "earlier": "Partner recovery workbook (earlier round)",
                "reinstate": "MSNA team - duration-rule reinstatement of the first interview at this identifier (2026-10-02)"}
    src = {c["sources"]: "; ".join(dict.fromkeys(src_part.get(p, p) for p in c["sources"].split("+")))}
    return dict(ident(u), check_binding=f"{var} ~/~ {u}", check_id=cid, issue=issue, label=var, question=var,
                old_value=eff_old, change_type="change_response", new_value=new, Data_Feedback=feedback,
                FO_Comments=f"Changed from {eff_old or 'blank'} to {new}. {c['justification']}",
                response_file=src.get(c["sources"], c["sources"]), check_date=(t.get("detected_date") or "2026-10-02"),
                decision_basis=basis, decided_by=c["decided_by"], review_reminder=c["reminder"], already_in_do_master_log=in_do)


def write(out, sheet_rows, stem, preview, r1members):
    if preview:
        stem = stem.with_name(stem.name + "_PREVIEW")
    with open(f"{stem}.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=DO_COLS + EXTRA)
        w.writeheader()
        w.writerows(out)
    wbo = openpyxl.Workbook()
    ws = wbo.active
    ws.title = "cleaning_log"
    ws.append(DO_COLS + EXTRA)
    for o in out:
        ws.append([o[c] for c in DO_COLS + EXTRA])
    for i, c in enumerate(ws[1], start=1):
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="7F6000" if i > len(DO_COLS) else "1F4E78")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    sm = wbo.create_sheet("summary")
    sm.append(["change_type", "check_id", "rows", "interviews"])
    grp, cnt = defaultdict(set), Counter()
    for o in out:
        grp[(o["change_type"], o["check_id"])].add(o["uuid"])
        cnt[(o["change_type"], o["check_id"])] += 1
    for (ct, cid), n in sorted(cnt.items()):
        sm.append([ct, cid, n, len(grp[(ct, cid)])])
    sm.append([])
    sm.append(["org_id", "interviews removed", "interviews corrected (kept)", "interviews flagged, kept unchanged"])
    by_org = defaultdict(lambda: [set(), set(), set()])
    for o in out:
        by_org[o["org_id"] or "(not recorded)"][{"remove_survey": 0, "change_response": 1, "no_action": 2}[o["change_type"]]].add(o["uuid"])
    for k in sorted(by_org):
        a, b, c = by_org[k]
        sm.append([k, len(a), len(b - a), len(c - a - b)])
    sm.append([])
    sm.append(["removals by round", "deletion_round", "interview_round", "interviews removed"])
    for (dr, ir), n in sorted(Counter((o["deletion_round"], o["interview_round"]) for o in out if o["change_type"] == "remove_survey").items()):
        sm.append(["", dr, ir, n])
    for col, wdt in zip("ABCD", (26, 34, 26, 34)):
        sm.column_dimensions[col].width = wdt

    final_rule_rows = Counter(o["decision_basis"].split(":")[0] for o in out if "applied to the final dataset" in o["decision_basis"])
    rm = wbo.create_sheet("readme")
    lines = [
        ["FINAL DELETION & CORRECTION LOG - MSNA N-WEC 2026", ""],
        [f"Built {TODAY} by the MSNA monitoring team, for incorporation into the master cleaning log. It replaces the Round 1 "
         f"log of 2 Oct 2026 (MSNA_N-WEC_2026_Round1_deletion_log_2026-10-02) - every Round 1 row is carried over here.", ""],
        ["Scope", "Every submission in the final dataset (data collection closed 15 Oct 2026; final download 16 Oct 2026), "
                  f"Round 1 ({len(r1members):,} submissions up to 1 Oct) and Round 2 (everything after). Every recovery item is "
                  "resolved - nothing is left pending."],
        ["deletion_round", "Removals only. Round 1 = already removed in the Round 1 log sent on 2 Oct. Round 2 = removed since: "
                           "interviews collected after Round 1, and Round 1 interviews removed after the Round 1 submission (for "
                           "example the short interviews whose audit files arrived after 2 Oct)."],
        ["interview_round", "Every row. Round 1 = one of the 28,046 Round 1 submissions; Round 2 = submitted after them."],
        ["Format", "The master cleaning log's 29 columns, change_type and status values; every row is status = addressed. Seven traceability columns follow (orange headers)."],
        ["remove_survey", "Delete the survey. One row per removed interview, labelled by the real basis of the removal (consent first, then the duration rule, then the rest - the master log's own priority)."],
        ["change_response", "Set the variable in 'question' to new_value. One row per variable - every variable the data officer's checks read is set, e.g. both non_idp_point_id and sample_point_id."],
        ["no_action", "Flagged, reviewed, interview kept unchanged. Also one row per interview in a cluster that never had a household listing (the master log's own listing_duplicate_draw pattern)."],
        ["Justification", "issue = what was found; Data_Feedback = the evidence; FO_Comments = the decision and why; response_file = where the decision came from."],
        ["decision_basis", "The rule or source the decision rests on: a validated threshold, a partner's response, or a named closeout rule (Jack, 1-2 Oct 2026; the same rules were applied to the final dataset on 16 Oct 2026, and those rows say so)."],
        ["decided_by", "partner = confirmed by the partner in a recovery workbook; internal_team = decided by the MSNA team under the stated rule."],
        ["tracker_issue_id", "The row's id(s) in the MSNA team's recovery issue tracker, for tracing its full history."],
        ["review_reminder", "Filled where the MSNA team asks for a check before the interview is used: GPS-placed interviews (population group never recorded), interviews kept 500 m-2 km from the nearest drawn point, partner answers the device GPS does not support, household numbers outside the latest listing, and one removal the audit trail no longer supports."],
        ["already_in_do_master_log", "Compared with the master cleaning log of 1 Oct 2026, the latest the MSNA team holds - so it says 'no' for every later interview. yes = the master log already does the same: don't apply it twice. follows = the master log first recovers this variable (auto_recovery_sampling_fields), and this row applies after it - old_value is the recovered value. CONFLICT = the master log does something different and this log supersedes it (the interviews the master log removes for sampling_assignment_missing that this log keeps, placed by device GPS; and one removal, for a reason other than duration, of an interview the master log keeps as duration_low)."],
        ["", ""],
        ["Suffixes _b, _c ...", "A second (third ...) DIFFERENT household recorded at one drawn sample point or listed household number (household size or head's sex/age differ from the first): both are kept, and the later one's identifier is suffixed so key-based duplicate checks don't re-flag it. Cluster and stratum are unchanged."],
        ["Integer questions", "idp_hh_number_from_listing and idp_walk_position are integer questions in the tool: a suffix makes the value text. gps_cluster_checks.R builds its duplicate key with dplyr::coalesce() of these two fields, which needs both columns to be the same type - convert both to character before running it."],
        ["Duration rule", "An interview is removed for short duration only when its audit-trail duration (cleaningtools' create_duration_from_audit_sum_all - the time actually spent on the questions), rounded to one decimal place of a minute, is below 20.0 minutes - that is, when it is under 19.95 minutes. This is the same rounding the master log's own duration check uses (decided by the MSNA team lead, 2 Oct 2026). Durations were recomputed by the MSNA team on the full audit export."],
        ["Incomplete audit trails", "A few audit trails have a question with no end or start time (typically the app was "
                                    "interrupted), so the duration formula cannot total them. Each such interview is kept when "
                                    "the questions with both times recorded already add up to 19.95 minutes or more - the "
                                    "missing times can only add to it (MSNA team lead, 16 Oct 2026). no_action rows, check_id "
                                    "duration_audit_incomplete, state each interview's minimum."],
        ["Reinstated interviews", "Five interviews removed earlier by other processes do not meet this rule and are reinstated (no_action rows, check_id duration_under_20): four of 19.96-19.99 minutes (20.0 at one decimal place) and one of 24.4 minutes whose earlier flag the full audit trail does not support. Three of them share their sample point / listed household with a later interview of a different household (a re-visit after the removal); the earlier upload keeps it and the later interview is reassigned or suffixed - change_response rows."],
        ["GPS distances", "dist_btn_sample_collected is recomputed (device GPS to the new point, the form's own formula) only where an interview's point actually moved and its device GPS was in the raw export of 29 Sep."],
    ] + ([["Final closeout (16 Oct)", "Flags still open when data collection closed were settled by the same closeout rules, "
           "applied to the final dataset: " + "; ".join(f"{k} ({n} rows)" for k, n in sorted(final_rule_rows.items())) +
           ". No identifier was changed in the final round."]] if final_rule_rows else []) + [
        ["", ""],
        ["Master cleaning log readme (for reference):", ""],
    ] + sheet_rows["readme"]
    for line in lines:
        rm.append(line)
    rm.column_dimensions["A"].width = 30
    rm.column_dimensions["B"].width = 140
    for row in rm.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    vr = wbo.create_sheet("validation_rules")
    for line in sheet_rows["validation_rules"]:
        vr.append(line)
    wbo.save(f"{stem}.xlsx")

    print(f"wrote {stem.name}.xlsx/.csv: {len(out)} rows, {len({o['uuid'] for o in out})} interviews")
    print("interviews by change_type:", {k: len({o['uuid'] for o in out if o['change_type'] == k}) for k in ("remove_survey", "change_response", "no_action")})
    for k, n in sorted(Counter((o["change_type"], o["check_id"]) for o in out).items()):
        print(f"  {n:5d}  {k[0]:16s} {k[1]}")
    print("review reminders:", sum(1 for o in out if o["review_reminder"]), "rows /", len({o["uuid"] for o in out if o["review_reminder"]}), "interviews")
    print("already in master log:", dict(Counter(o["already_in_do_master_log"].split(" - ")[0] for o in out)))
    print("decided_by:", dict(Counter(o["decided_by"] for o in out)))


if __name__ == "__main__":
    main()
